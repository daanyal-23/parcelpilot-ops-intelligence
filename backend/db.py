import sqlite3
import os
import openpyxl
from datetime import datetime
import pytz
import uuid

DATASET_SNAPSHOT_TIME = "2026-08-16T11:00:00+05:30"
IST = pytz.timezone("Asia/Kolkata")
DB_PATH = os.path.join(os.path.dirname(__file__), "parcelpilot.db")
DATA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data"))
EXCEL_PATH = os.path.join(DATA_DIR, "ParcelPilot_Assessment_Data.xlsx")

def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("DROP TABLE IF EXISTS accounts")
    cursor.execute("DROP TABLE IF EXISTS orders")
    cursor.execute("DROP TABLE IF EXISTS tickets")
    cursor.execute("DROP TABLE IF EXISTS actions")
    cursor.execute("DROP TABLE IF EXISTS issued_credits")
    cursor.execute("DROP TABLE IF EXISTS user_sessions")

    cursor.execute("""
    CREATE TABLE accounts (
        account_id TEXT PRIMARY KEY,
        account_name TEXT NOT NULL,
        plan TEXT NOT NULL,
        status TEXT NOT NULL,
        csm TEXT,
        contract_file TEXT,
        premium_support BOOLEAN,
        notes TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE orders (
        order_id TEXT PRIMARY KEY,
        account_id TEXT NOT NULL,
        carrier TEXT NOT NULL,
        status TEXT NOT NULL,
        booked_at TEXT NOT NULL,
        pickup_window_start TEXT,
        pickup_window_end TEXT,
        pickup_actual_at TEXT,
        shipment_fee_inr REAL NOT NULL,
        carrier_fault BOOLEAN,
        customer_fault BOOLEAN,
        cancellation_requested_at TEXT,
        notes TEXT,
        FOREIGN KEY (account_id) REFERENCES accounts(account_id)
    )
    """)

    cursor.execute("""
    CREATE TABLE tickets (
        ticket_id TEXT PRIMARY KEY,
        account_id TEXT NOT NULL,
        created_at TEXT NOT NULL,
        status TEXT NOT NULL,
        subject TEXT NOT NULL,
        description TEXT NOT NULL,
        channel TEXT NOT NULL,
        assigned_to TEXT,
        last_customer_message_at TEXT,
        historical_resolution TEXT,
        FOREIGN KEY (account_id) REFERENCES accounts(account_id)
    )
    """)

    cursor.execute("""
    CREATE TABLE actions (
        action_id TEXT PRIMARY KEY,
        action_type TEXT NOT NULL,
        target_id TEXT NOT NULL,
        payload TEXT NOT NULL,
        status TEXT NOT NULL, -- pending_confirmation, confirmed, executed, cancelled, rejected
        created_by_role TEXT NOT NULL,
        summary TEXT NOT NULL,
        created_at TEXT NOT NULL,
        executed_at TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE issued_credits (
        credit_id TEXT PRIMARY KEY,
        action_id TEXT NOT NULL,
        order_id TEXT NOT NULL,
        account_id TEXT NOT NULL,
        amount_inr REAL NOT NULL,
        approved_by_user TEXT NOT NULL,
        approved_by_role TEXT NOT NULL,
        issued_at TEXT NOT NULL,
        notes TEXT,
        FOREIGN KEY (action_id) REFERENCES actions(action_id),
        FOREIGN KEY (order_id) REFERENCES orders(order_id),
        FOREIGN KEY (account_id) REFERENCES accounts(account_id)
    )
    """)

    cursor.execute("""
    CREATE TABLE user_sessions (
        session_id TEXT PRIMARY KEY,
        user_id TEXT NOT NULL,
        role TEXT NOT NULL,
        created_at TEXT NOT NULL
    )
    """)

    conn.commit()
    conn.close()

def create_session(user_id: str, role: str) -> str:
    session_id = f"sess_{uuid.uuid4().hex[:16]}"
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO user_sessions (session_id, user_id, role, created_at) VALUES (?, ?, ?, ?)",
            (session_id, user_id, role, DATASET_SNAPSHOT_TIME)
        )
        conn.commit()
    finally:
        conn.close()
    return session_id

def get_session(session_id: str):
    if not session_id:
        return None
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM user_sessions WHERE session_id = ?", (session_id,))
        row = cursor.fetchone()
        return dict(row) if row else None
    finally:
        conn.close()

def format_dt(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        if val.tzinfo is None:
            val = IST.localize(val)
        return val.isoformat()
    if isinstance(val, str):
        val = val.strip()
        if not val:
            return None
        # Parse 'YYYY-MM-DD HH:MM' or ISO format
        try:
            dt = datetime.strptime(val, "%Y-%m-%d %H:%M")
            return IST.localize(dt).isoformat()
        except ValueError:
            try:
                dt = datetime.fromisoformat(val)
                if dt.tzinfo is None:
                    dt = IST.localize(dt)
                return dt.isoformat()
            except ValueError:
                return val
    return str(val)

def seed_db():
    init_db()
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Excel dataset not found at {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    conn = get_db_connection()
    cursor = conn.cursor()

    # Seed Accounts
    ws_acc = wb["accounts"]
    rows = list(ws_acc.iter_rows(values_only=True))
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        cursor.execute(
            """
            INSERT INTO accounts (account_id, account_name, plan, status, csm, contract_file, premium_support, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (r[0], r[1], r[2], r[3], r[4], r[5], bool(r[6]), r[7])
        )

    # Seed Orders
    ws_ord = wb["orders"]
    rows = list(ws_ord.iter_rows(values_only=True))
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        cursor.execute(
            """
            INSERT INTO orders (order_id, account_id, carrier, status, booked_at, pickup_window_start,
                               pickup_window_end, pickup_actual_at, shipment_fee_inr, carrier_fault,
                               customer_fault, cancellation_requested_at, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                r[0], r[1], r[2], r[3],
                format_dt(r[4]),
                format_dt(r[5]),
                format_dt(r[6]),
                format_dt(r[7]),
                float(r[8]) if r[8] is not None else 0.0,
                bool(r[9]) if r[9] is not None else False,
                bool(r[10]) if r[10] is not None else False,
                format_dt(r[11]),
                r[12]
            )
        )

    # Seed Tickets
    ws_tkt = wb["tickets"]
    rows = list(ws_tkt.iter_rows(values_only=True))
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        cursor.execute(
            """
            INSERT INTO tickets (ticket_id, account_id, created_at, status, subject, description,
                                channel, assigned_to, last_customer_message_at, historical_resolution)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                r[0], r[1],
                format_dt(r[2]),
                r[3], r[4], r[5], r[6], r[7],
                format_dt(r[8]),
                r[9]
            )
        )

    # Create default test sessions
    cursor.execute(
        "INSERT INTO user_sessions (session_id, user_id, role, created_at) VALUES (?, ?, ?, ?)",
        ("sess_default_agent", "agent_rohit", "support_agent", DATASET_SNAPSHOT_TIME)
    )
    cursor.execute(
        "INSERT INTO user_sessions (session_id, user_id, role, created_at) VALUES (?, ?, ?, ?)",
        ("sess_default_manager", "mgr_priya", "manager", DATASET_SNAPSHOT_TIME)
    )

    conn.commit()
    conn.close()
    print(f"Database seeded successfully at {DB_PATH}")

if __name__ == "__main__":
    seed_db()
